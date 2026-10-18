import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from fastapi.testclient import TestClient
import pytest
import pymupdf
import docx

from redact_core.api import app
from scripts.generate_fixtures import FIXTURE_PII

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"


@pytest.fixture
def client():
    return TestClient(app)


def test_api_health(client):
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "version" in data
    assert "sales-quarterly" in data["policies"]


def test_api_policies_and_subjects(client):
    p_resp = client.get("/api/policies")
    assert p_resp.status_code == 200
    policies = p_resp.json()
    policy_ids = [p["policy_id"] for p in policies]
    assert "sales-quarterly" in policy_ids
    assert "finance-full-access" in policy_ids
    assert "reviewer" in policy_ids

    s_resp = client.get("/api/subjects")
    assert s_resp.status_code == 200
    subjects = s_resp.json()
    subj_ids = [s["subject_id"] for s in subjects]
    assert "demo_sales_viewer" in subj_ids


def test_api_upload_and_create_view_pdf(client):
    pdf_path = FIXTURES_DIR / "sample_report.pdf"
    assert pdf_path.exists()

    with open(pdf_path, "rb") as f:
        files = {"file": ("sample_report.pdf", f, "application/pdf")}
        up_resp = client.post("/api/upload", files=files)
    assert up_resp.status_code == 200
    up_data = up_resp.json()
    assert up_data["file_type"] == "pdf"
    assert up_data["page_count"] == 2
    assert up_data["findings_count"] > 0
    source_sha256 = up_data["source_sha256"]

    # Generate view
    view_req = {
        "source_sha256": source_sha256,
        "ext": "pdf",
        "policy_id": "sales-quarterly",
        "subject_id": "demo_sales_viewer",
        "roles": ["sales", "viewer"],
    }
    v_resp = client.post("/api/views", json=view_req)
    assert v_resp.status_code == 200
    v_data = v_resp.json()
    assert v_data["verification"]["passed"] is True
    assert "view_id" in v_data
    view_id = v_data["view_id"]

    # Check manifest
    m_resp = client.get(f"/api/views/{view_id}/manifest")
    assert m_resp.status_code == 200
    m_data = m_resp.json()
    assert m_data["view_id"] == view_id
    assert m_data["source_sha256"] == source_sha256

    # Download and verify sensitive PII absence
    dl_resp = client.get(f"/api/views/{view_id}/download")
    assert dl_resp.status_code == 200
    derived_bytes = dl_resp.content

    doc = pymupdf.open(stream=derived_bytes, filetype="pdf")
    all_text = " ".join(page.get_text() for page in doc)
    doc.close()

    for pii_key in ["salary_1", "salary_2", "emp_id_1", "emp_id_2", "phone_1", "phone_2"]:
        val = FIXTURE_PII[pii_key]
        assert val not in all_text, f"PII leak detected: {pii_key} ('{val}') found in PDF output"


def test_api_upload_and_create_view_docx(client):
    docx_path = FIXTURES_DIR / "sample_report.docx"
    assert docx_path.exists()

    with open(docx_path, "rb") as f:
        files = {"file": ("sample_report.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        up_resp = client.post("/api/upload", files=files)
    assert up_resp.status_code == 200
    up_data = up_resp.json()
    assert up_data["file_type"] == "docx"
    assert up_data["findings_count"] > 0
    source_sha256 = up_data["source_sha256"]

    # Generate view
    view_req = {
        "source_sha256": source_sha256,
        "ext": "docx",
        "policy_id": "sales-quarterly",
        "subject_id": "demo_sales_viewer",
        "roles": ["sales", "viewer"],
    }
    v_resp = client.post("/api/views", json=view_req)
    assert v_resp.status_code == 200
    v_data = v_resp.json()
    assert v_data["verification"]["passed"] is True
    view_id = v_data["view_id"]

    # Download and verify
    dl_resp = client.get(f"/api/views/{view_id}/download")
    assert dl_resp.status_code == 200
    from io import BytesIO
    d = docx.Document(BytesIO(dl_resp.content))
    all_text = " ".join(p.text for p in d.paragraphs) + " " + " ".join(
        c.text for t in d.tables for r in t.rows for c in r.cells
    )

    for pii_key in ["salary_1", "salary_2", "emp_id_1", "emp_id_2", "phone_1", "phone_2"]:
        val = FIXTURE_PII[pii_key]
        assert val not in all_text, f"PII leak detected: {pii_key} ('{val}') found in DOCX output"


def test_api_scanned_fail_closed(client):
    scanned_path = FIXTURES_DIR / "scanned_sample.pdf"
    assert scanned_path.exists()

    with open(scanned_path, "rb") as f:
        files = {"file": ("scanned_sample.pdf", f, "application/pdf")}
        up_resp = client.post("/api/upload", files=files)
    assert up_resp.status_code == 200
    up_data = up_resp.json()
    assert up_data["is_scanned"] is True
    source_sha256 = up_data["source_sha256"]

    # Generation should be blocked fail-closed (HTTP 403)
    view_req = {
        "source_sha256": source_sha256,
        "ext": "pdf",
        "policy_id": "sales-quarterly",
        "subject_id": "demo_sales_viewer",
        "roles": ["sales", "viewer"],
    }
    v_resp = client.post("/api/views", json=view_req)
    assert v_resp.status_code == 403
    assert "blocked" in v_resp.json()["detail"].lower()


def test_api_healthz(client):
    resp = client.get("/healthz")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_api_health_storage_metadata(client):
    resp = client.get("/health")
    assert resp.status_code == 200
    data = resp.json()
    assert "storage_root" in data
    assert "storage_ephemeral" in data
    assert isinstance(data["storage_ephemeral"], bool)


def test_cors_behavior(client):
    # Allowed origin receives CORS headers
    resp = client.get("/healthz", headers={"Origin": "http://localhost:8080"})
    assert resp.status_code == 200
    assert resp.headers.get("access-control-allow-origin") == "http://localhost:8080"
    assert resp.headers.get("access-control-allow-credentials") == "true"

    # Unlisted origin is not reflected in access-control-allow-origin
    unauthorized = client.get("/healthz", headers={"Origin": "https://unauthorized-domain.example.com"})
    assert unauthorized.status_code == 200
    assert unauthorized.headers.get("access-control-allow-origin") is None


def test_cors_origins_helper(monkeypatch):
    import os
    from redact_core.api import get_cors_origins

    monkeypatch.setenv("REDACTCORE_CORS_ORIGINS", "https://app.argus.internal, https://demo.argus.internal")
    origins = get_cors_origins()
    assert origins == ["https://app.argus.internal", "https://demo.argus.internal"]

    monkeypatch.delenv("REDACTCORE_CORS_ORIGINS", raising=False)
    default_origins = get_cors_origins()
    assert "http://localhost:8080" in default_origins
    assert "*" not in default_origins


def test_safe_filename_handling_and_path_traversal(client):
    # Path traversal patterns must be rejected at API boundary
    traversal_filenames = [
        "../../etc/passwd.pdf",
        "..\\..\\windows\\win.ini.pdf",
        "/tmp/absolute_path.pdf",
        "subfolder/test.pdf",
        "test\x00traversal.pdf",
        "   ",
        "invalid_ext.exe",
    ]
    for bad_name in traversal_filenames:
        resp = client.post(
            "/api/upload",
            files={"file": (bad_name, b"dummy content", "application/pdf")},
        )
        assert resp.status_code == 400, f"Expected 400 for '{bad_name}', got {resp.status_code}"


def test_upload_empty_file_rejected(client):
    resp = client.post(
        "/api/upload",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()


def test_upload_max_size_enforcement(client, monkeypatch):
    # Enforce boundary max upload bytes
    monkeypatch.setenv("REDACTCORE_MAX_UPLOAD_BYTES", "50")
    large_payload = b"%PDF-1.4 " + b"X" * 100
    resp = client.post(
        "/api/upload",
        files={"file": ("large.pdf", large_payload, "application/pdf")},
    )
    assert resp.status_code == 413
    assert "exceeds maximum allowed upload size" in resp.json()["detail"].lower()


def test_fixture_gating_and_traversal_rejection(client, monkeypatch):
    import os

    # When disabled (default in hosted mode), returns 403 Forbidden
    monkeypatch.delenv("REDACTCORE_ENABLE_DEMO_FIXTURES", raising=False)
    resp = client.get("/api/fixtures/sample_report.pdf")
    assert resp.status_code == 403
    assert "disabled" in resp.json()["detail"].lower()

    # When enabled, serves synthetic fixture
    monkeypatch.setenv("REDACTCORE_ENABLE_DEMO_FIXTURES", "true")
    resp_enabled = client.get("/api/fixtures/sample_report.pdf")
    assert resp_enabled.status_code == 200

    # Path traversal within fixture route must be rejected
    resp_traversal = client.get("/api/fixtures/..%2F..%2Fpyproject.toml")
    assert resp_traversal.status_code in (400, 404)


def test_view_id_validation_path_traversal(client):
    # Invalid characters or traversal in view_id rejected with 400
    resp1 = client.get("/api/views/../../etc/passwd/manifest")
    assert resp1.status_code in (400, 404)

    resp2 = client.get("/api/views/invalid;id/download")
    assert resp2.status_code == 400


def test_generate_view_request_validation(client):
    # Invalid sha256 (not 64-char hex)
    resp = client.post("/api/views", json={
        "source_sha256": "invalid_hash",
        "ext": "pdf",
        "policy_id": "sales-quarterly",
        "subject_id": "demo_sales_viewer",
    })
    assert resp.status_code == 422

    # Invalid extension
    resp_ext = client.post("/api/views", json={
        "source_sha256": "a" * 64,
        "ext": "exe",
        "policy_id": "sales-quarterly",
        "subject_id": "demo_sales_viewer",
    })
    assert resp_ext.status_code == 422


def test_activate_policy(client):
    policy_payload = {
        "policy_id": "test-dynamic-policy",
        "name": "Dynamic Test Policy",
        "version": "1.0.0",
        "description": "Test policy committed dynamically",
        "default_action": "allow",
        "require_figure_review": False,
        "rules": [
            {
                "entity_type": "SALARY",
                "action": "redact",
                "replacement": "[CUSTOM SALARY REDACTED]",
                "description": "Redact compensation",
            }
        ],
    }
    resp = client.post("/api/policies", json=policy_payload)
    assert resp.status_code == 201
    data = resp.json()
    assert data["status"] == "active"
    assert data["policy_id"] == "test-dynamic-policy"

    # Verify it shows up in GET /api/policies
    list_resp = client.get("/api/policies")
    assert list_resp.status_code == 200
    ids = [p["policy_id"] for p in list_resp.json()]
    assert "test-dynamic-policy" in ids


def test_cannot_delete_reserved_demo_policy(client):
    """Verifies that system-reserved demo policies cannot be deleted (HTTP 403 Forbidden)."""
    reserved_demos = [
        "engineering-internal-strict",
        "creative-partner-minimal",
        "thirdparty-vendor-restricted",
        "public-relations-safe",
        "sales-quarterly",
        "finance-full-access",
        "reviewer",
        "market-full-access",
    ]
    for pol_id in reserved_demos:
        del_resp = client.delete(f"/api/policies/{pol_id}")
        assert del_resp.status_code == 403
        data = del_resp.json()
        assert "cannot be deleted" in data["detail"].lower()


def test_delete_custom_policy(client):
    """Verifies that a custom policy can be created and cleanly deleted."""
    custom_policy = {
        "policy_id": "test-to-delete",
        "name": "Custom Ephemeral Policy",
        "version": "1.0.0",
        "description": "Ephemeral policy to test deletion workflow",
        "default_action": "allow",
        "require_figure_review": False,
        "rules": [
            {
                "entity_type": "SALARY",
                "action": "redact",
                "replacement": "[CUSTOM DELETED]",
                "description": "Redact compensation",
            }
        ],
    }
    # 1. Create custom policy
    create_resp = client.post("/api/policies", json=custom_policy)
    assert create_resp.status_code == 201

    # Verify present
    list_resp = client.get("/api/policies")
    assert any(p["policy_id"] == "test-to-delete" for p in list_resp.json())

    # 2. Delete custom policy
    del_resp = client.delete("/api/policies/test-to-delete")
    assert del_resp.status_code == 200
    del_data = del_resp.json()
    assert del_data["status"] == "deleted"
    assert del_data["policy_id"] == "test-to-delete"

    # 3. Verify removed from list
    list_resp_after = client.get("/api/policies")
    assert not any(p["policy_id"] == "test-to-delete" for p in list_resp_after.json())

    # 4. Deleting again returns 404
    del_again = client.delete("/api/policies/test-to-delete")
    assert del_again.status_code == 404

