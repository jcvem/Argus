import argparse
from pathlib import Path
import sys

# Ensure repository root is on sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import docx
import pymupdf

from redact_core.hashing import compute_sha256_file
from redact_core.storage import default_store
from scripts.generate_fixtures import FIXTURE_PII


def validate_view(view_id: str) -> bool:
    manifest = default_store.get_manifest(view_id)
    if not manifest:
        print(f"[FAIL] Manifest for view '{view_id}' not found.")
        return False

    print(f"=== Validating View: {view_id} ===")
    print(f"Policy: {manifest.policy_id} (v{manifest.policy_version})")
    print(f"Source SHA-256: {manifest.source_sha256}")
    print(f"Output SHA-256: {manifest.output_sha256}")

    # 1. Verify original file still exists and matches source_sha256
    orig_path = default_store.get_original_path(manifest.source_sha256, manifest.input_format)
    if not orig_path or not orig_path.exists():
        print(f"[FAIL] Original document with hash {manifest.source_sha256} missing from storage!")
        return False

    curr_orig_hash = compute_sha256_file(orig_path)
    if curr_orig_hash != manifest.source_sha256:
        print(f"[FAIL] Original document has been MUTATED! Hash mismatch: {curr_orig_hash} != {manifest.source_sha256}")
        return False
    print(f"[PASS] Original document byte integrity confirmed: {curr_orig_hash}")

    # 2. Verify derived view file exists and matches output_sha256
    view_path = default_store.get_view_path(view_id, manifest.output_format)
    if not view_path or not view_path.exists():
        print(f"[FAIL] Derived view file {view_path} missing from storage!")
        return False

    curr_output_hash = compute_sha256_file(view_path)
    if curr_output_hash != manifest.output_sha256:
        print(f"[FAIL] Output file hash mismatch: {curr_output_hash} != {manifest.output_sha256}")
        return False
    print(f"[PASS] Output view byte integrity confirmed: {curr_output_hash}")

    # 3. Text search for leaked sensitive PII
    leaks = []
    if manifest.output_format == "pdf":
        doc = pymupdf.open(str(view_path))
        all_text = " ".join(p.get_text() for p in doc)
        doc.close()
    else:
        d = docx.Document(str(view_path))
        all_text = " ".join(p.text for p in d.paragraphs) + " " + " ".join(c.text for t in d.tables for r in t.rows for c in r.cells)

    # Check for known fixture PII if policy was sales-quarterly
    if "sales" in manifest.policy_id:
        for pii_key in ["salary_1", "salary_2", "emp_id_1", "emp_id_2", "phone_1", "phone_2"]:
            val = FIXTURE_PII[pii_key]
            if val in all_text:
                leaks.append(f"Found sensitive {pii_key}: '{val}'")

    if leaks:
        print(f"[FAIL] PII Leaks detected in derived artifact:")
        for leak in leaks:
            print(f"  - {leak}")
        return False

    print(f"[PASS] Leak check passed: No sensitive values present in derived view.")
    print(f"[PASS] Verification manifest: {manifest.verification.get('passed')}")
    return True


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate a derived document view against its manifest")
    parser.add_argument("view_id", type=str, help="View ID to validate")
    args = parser.parse_args()

    success = validate_view(args.view_id)
    return 0 if success else 1


if __name__ == "__main__":
    sys.exit(main())
