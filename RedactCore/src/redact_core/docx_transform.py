"""DOCX derived-view transformation engine.

Applies redactions, replacements, and masking to a copy of a DOCX document,
cleans core metadata properties, and preserves document structure.
"""

from io import BytesIO
from pathlib import Path
import re
import docx

from redact_core.models import RedactionOperation, RedactionPlan


def _replace_text_in_paragraph(paragraph: docx.text.paragraph.Paragraph, target_text: str, replacement_text: str) -> bool:
    """Replaces occurrences of target_text within paragraph runs."""
    if target_text not in paragraph.text:
        return False

    # First attempt: exact match in individual runs (preserves run formatting)
    replaced_in_runs = False
    for run in paragraph.runs:
        if target_text in run.text:
            run.text = run.text.replace(target_text, replacement_text)
            replaced_in_runs = True

    # If replaced in individual runs and no more occurrences exist, we're done
    if replaced_in_runs and target_text not in paragraph.text:
        return True

    # Second attempt: cross-run boundary match
    # If the text spans across runs, replace in the full paragraph text and consolidate
    if target_text in paragraph.text:
        full_text = paragraph.text.replace(target_text, replacement_text)
        # Clear all runs and put replaced text into first run
        if paragraph.runs:
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = full_text
        return True

    return False


def _replace_text_in_table(table: docx.table.Table, target_text: str, replacement_text: str) -> bool:
    """Replaces occurrences of target_text inside table cells."""
    changed = False
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if _replace_text_in_paragraph(p, target_text, replacement_text):
                    changed = True
    return changed


def transform_docx(
    original_bytes: bytes,
    plan: RedactionPlan,
    sanitize_metadata: bool = True,
) -> bytes:
    """Transforms a copy of the original DOCX document according to the redaction plan."""
    if plan.blocked:
        raise ValueError(f"Cannot transform document under a blocked plan: {plan.block_reason}")

    stream = BytesIO(original_bytes)
    doc = docx.Document(stream)

    for op in plan.operations:
        if op.action in ("allow",):
            continue

        target_text = op.target_text.strip()
        replacement = op.replacement_text or "[REDACTED]"
        if not target_text:
            continue

        # 1. Replace in main body paragraphs
        for p in doc.paragraphs:
            _replace_text_in_paragraph(p, target_text, replacement)

        # 2. Replace in tables
        for t in doc.tables:
            _replace_text_in_table(t, target_text, replacement)

        # 3. Replace in headers and footers
        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    _replace_text_in_paragraph(p, target_text, replacement)
            if section.footer:
                for p in section.footer.paragraphs:
                    _replace_text_in_paragraph(p, target_text, replacement)

    # Sanitize document properties
    if sanitize_metadata:
        cp = doc.core_properties
        cp.title = "Redacted View"
        cp.author = "Argus RedactCore"
        cp.comments = f"Derived privacy-preserved view under policy {plan.policy_id} (decision: {plan.decision_id})"
        cp.last_modified_by = "Argus RedactCore"

    out_stream = BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()
